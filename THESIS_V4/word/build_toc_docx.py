"""
Writes the table of contents of the thesis, with the page numbers of the built
Word thesis, as a separate Word file.

  python THESIS_V4/word/build_toc_docx.py <indexes.json> <out.docx>
"""
import json
import re
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from build_docx import AUTHOR, TITLE


def entries(text):
    out = []
    for line in text.split("\n"):
        line = line.strip("\t ")
        if not line or line == "Contents" or "\t" not in line:
            continue
        title, page = line.rsplit("\t", 1)
        out.append((title.replace("\t", " ").strip(), page.strip()))
    return out


def main(index_json, out):
    data = json.loads(open(index_json, encoding="utf-8").read())
    toc = next(x["text"] for x in data if x["name"].startswith("Table of Contents"))
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.59), Cm(27.94)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(2.5), Cm(1.5), Cm(2.5), Cm(2.5)
    text_w = s.page_width - s.left_margin - s.right_margin
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    st.paragraph_format.space_after = Pt(0)

    h = d.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.paragraph_format.space_after = Pt(18)
    r = h.add_run("Table of Contents")
    r.bold = True
    r.font.size = Pt(16)

    p = d.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(text_w, WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("\tPage")
    r.bold = True

    for title, page in entries(toc):
        if re.match(r"^\d+\.\d+\.\d+ ", title):
            num, rest, left, hang, bold, before = *title.split(" ", 1), Cm(1.9), Cm(1.4), False, 2
        elif re.match(r"^\d+\.\d+ ", title):
            num, rest, left, hang, bold, before = *title.split(" ", 1), Cm(0.9), Cm(1.0), False, 3
        else:
            num, rest, left, hang, bold, before = None, title, Cm(0), Cm(0), True, 9
        p = d.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = left
        pf.first_line_indent = -hang if num else None
        pf.space_before = Pt(before)
        pf.space_after = Pt(1)
        pf.keep_with_next = bold and title.startswith("Chapter")
        if num:
            pf.tab_stops.add_tab_stop(left, WD_TAB_ALIGNMENT.LEFT)
        pf.tab_stops.add_tab_stop(text_w, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        label = f"{num}\t{rest}" if num else rest
        r = p.add_run(f"{label}\t{page}")
        r.bold = bold

    cp = d.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.title = f"{TITLE} - Table of Contents"
    cp.comments = ""
    d.save(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
